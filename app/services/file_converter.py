"""
File format converter for E-Rakshak.
Converts XLSX, XLS, TSV files to CSV for pipeline processing.
"""

import os
import shutil
from pathlib import Path

import pandas as pd

from app.core.logging import logger


class FileConverter:
    """
    Converts tabular file formats to CSV for the ingestion pipeline.
    """

    SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".tsv"}
    # Conversion bomb guards: refuse to expand small files into giant frames.
    MAX_ROWS = 500_000
    MAX_COLS = 256
    MAX_CELLS = 5_000_000

    @staticmethod
    def is_supported(filename: str) -> bool:
        """Check if a file extension is supported."""
        ext = Path(filename).suffix.lower()
        return ext in FileConverter.SUPPORTED_EXTENSIONS

    @staticmethod
    def get_extension(filename: str) -> str:
        """Get the lowercase extension of a file."""
        return Path(filename).suffix.lower()

    @staticmethod
    def convert_to_csv(input_path: str, output_path: str) -> str:
        """
        Convert a supported file to CSV format.
        Returns the path to the converted CSV file.
        If the input is already CSV, it's copied as-is.
        """
        ext = Path(input_path).suffix.lower()

        if ext == ".csv":
            if os.path.exists(output_path):
                os.remove(output_path)
            shutil.copy(input_path, output_path)
            return output_path

        if ext in (".xlsx", ".xls"):
            # Guard against xlsx bomb: check dimensions before full expansion.
            try:
                import openpyxl

                wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
                try:
                    for ws in wb.worksheets:
                        if ws.max_row is not None and ws.max_row > FileConverter.MAX_ROWS:
                            raise ValueError(f"Excel sheet too large ({ws.max_row} rows).")
                        if ws.max_column is not None and ws.max_column > FileConverter.MAX_COLS:
                            raise ValueError(f"Excel sheet too wide ({ws.max_column} columns).")
                        if (ws.max_row or 0) * (ws.max_column or 0) > FileConverter.MAX_CELLS:
                            raise ValueError("Excel sheet exceeds cell limit.")
                        break  # first sheet gates the file
                finally:
                    try:
                        wb.close()
                    except Exception:
                        pass
            except ValueError:
                raise
            except Exception:
                pass  # fall through to pandas; dimension check is best-effort
            df = pd.read_excel(input_path, nrows=FileConverter.MAX_ROWS + 1)
            if len(df) > FileConverter.MAX_ROWS or df.shape[1] > FileConverter.MAX_COLS:
                raise ValueError("Excel file exceeds row/column limits.")
            df.to_csv(output_path, index=False)
            return output_path

        if ext == ".tsv":
            # Chunked read to avoid loading a TSV bomb fully into RAM.
            chunks = pd.read_csv(input_path, sep="\t", chunksize=100_000, low_memory=False)
            first = True
            total = 0
            for chunk in chunks:
                if chunk.shape[1] > FileConverter.MAX_COLS:
                    raise ValueError("TSV file too wide.")
                total += len(chunk)
                if total > FileConverter.MAX_ROWS:
                    raise ValueError("TSV file too large.")
                chunk.to_csv(output_path, index=False, mode="w" if first else "a", header=first)
                first = False
            return output_path

        raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def validate_and_convert(input_path: str, output_dir: str) -> tuple[str, str]:
        """
        Validate a file is supported and convert it to CSV.
        Returns (csv_path, original_extension).
        """
        ext = Path(input_path).suffix.lower()

        if ext not in FileConverter.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format '{ext}'. "
                f"Supported formats: {', '.join(FileConverter.SUPPORTED_EXTENSIONS)}"
            )

        filename = Path(input_path).stem
        output_path = os.path.join(output_dir, f"{filename}.csv")

        FileConverter.convert_to_csv(input_path, output_path)
        logger.info("file_converted", input=input_path, output=output_path, original_ext=ext)

        return output_path, ext
